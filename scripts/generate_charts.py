#!/usr/bin/env python3
"""ライフプランxlsxからシミュレーション結果を抜き出し、HTMLグラフを生成

使い方:
    python3 generate_charts.py --xlsx /path/to/lifeplan.xlsx --out-dir ~/charts
"""
import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("⚠️  openpyxl が必要です: pip3 install openpyxl")
    sys.exit(1)


def extract_asset_trajectory(xlsx_path: Path) -> dict:
    """ライフプランシートから年齢別の累計資産を抜き出す（キャッシュ値ベース）"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["ライフプランシート"]
    ages, assets = [], []
    for col in range(4, ws.max_column + 1):
        age = ws.cell(row=4, column=col).value
        asset = ws.cell(row=95, column=col).value
        if isinstance(age, (int, float)) and isinstance(asset, (int, float)):
            ages.append(int(age))
            assets.append(int(asset))
    return {"ages": ages, "assets": assets}


def render_html_chart(data: dict, out_path: Path) -> None:
    """資産推移をシンプルなHTMLグラフで出力"""
    if not data["ages"]:
        print("⚠️  数値データが見つかりません。Google Sheetsで開いて再保存してください。")
        return

    ages_json = ",".join(map(str, data["ages"]))
    assets_json = ",".join(map(str, data["assets"]))
    max_a = max(data["assets"]) if data["assets"] else 1

    html = f"""<!DOCTYPE html>
<html lang="ja"><head><meta charset="UTF-8"><title>資産推移</title>
<style>
  body {{ background: #fdf6f0; font-family: 'Hiragino Sans', sans-serif; padding: 40px; }}
  .card {{ background: white; padding: 40px; border-radius: 16px; max-width: 1100px; margin: 0 auto;
           box-shadow: 0 4px 20px rgba(0,0,0,0.08); border: 1.5px solid #fde8d8; }}
  h2 {{ color: #5c3d2e; text-align: center; margin-bottom: 24px; }}
  canvas {{ display: block; margin: 0 auto; }}
</style></head><body>
<div class="card">
  <h2>📈 累計資産の推移（年齢別）</h2>
  <canvas id="chart" width="1000" height="500"></canvas>
</div>
<script>
const ages = [{ages_json}];
const assets = [{assets_json}];
const maxA = {max_a};
const c = document.getElementById('chart').getContext('2d');
const W = 1000, H = 500, ML = 80, MR = 30, MT = 30, MB = 60;
const PW = W - ML - MR, PH = H - MT - MB;
// 軸
c.strokeStyle = '#aaa'; c.beginPath();
c.moveTo(ML, MT); c.lineTo(ML, H-MB); c.lineTo(W-MR, H-MB); c.stroke();
// 線
c.strokeStyle = '#e07b39'; c.lineWidth = 3; c.beginPath();
ages.forEach((a, i) => {{
  const x = ML + (i / (ages.length-1)) * PW;
  const y = H - MB - (assets[i] / maxA) * PH;
  if (i === 0) c.moveTo(x, y); else c.lineTo(x, y);
}});
c.stroke();
// ラベル
c.fillStyle = '#5c3d2e'; c.font = '14px sans-serif';
[0, Math.floor(ages.length/4), Math.floor(ages.length/2), Math.floor(3*ages.length/4), ages.length-1].forEach(i => {{
  const x = ML + (i / (ages.length-1)) * PW;
  c.fillText(ages[i] + '歳', x - 12, H - MB + 20);
}});
// 縦軸
for (let i = 0; i <= 4; i++) {{
  const v = (maxA / 4) * i;
  const y = H - MB - (v / maxA) * PH;
  c.fillText((v/10000).toFixed(0) + '万', 10, y + 4);
}}
</script>
</body></html>"""
    out_path.write_text(html, encoding="utf-8")
    print(f"✓ チャート生成: {out_path}")


def main():
    parser = argparse.ArgumentParser(description="ライフプランの可視化")
    parser.add_argument("--xlsx", type=Path, required=True, help="ライフプランxlsxパス")
    parser.add_argument("--out-dir", type=Path, default=Path.home() / "lifeplan_charts",
                        help="出力先ディレクトリ")
    args = parser.parse_args()

    if not args.xlsx.exists():
        print(f"❌ xlsxが見つかりません: {args.xlsx}")
        sys.exit(1)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    data = extract_asset_trajectory(args.xlsx)
    render_html_chart(data, args.out_dir / "asset_trajectory.html")
    print(f"\n📊 ブラウザで開く: open {args.out_dir / 'asset_trajectory.html'}")


if __name__ == "__main__":
    main()
